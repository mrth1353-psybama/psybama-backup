import io
import os
from flask import Blueprint, request, jsonify, session, send_file, render_template
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from auth import admin_required
from models import db, User, Conversation, Message, Assessment, ContactRequest, AssessmentLead, Order, Product, WaaqAssessment
from models import CareerKnotAssessment
from models import WebinarRegistration
from models import CareerIntake, OrgIntake
from models import iran_now

KNOT_WEBINAR_TITLE = 'ریشه‌یابی گره کور شغلی'

admin_bp = Blueprint('admin', __name__)


@admin_bp.route('/')
@admin_bp.route('')
def admin_index():
    return render_template('admin.html')


@admin_bp.route('/login', methods=['POST'])
def admin_login():
    data = request.get_json()
    if not data or not data.get('password'):
        return jsonify({'error': 'password_required'}), 400

    if data['password'] == os.getenv('ADMIN_PASSWORD', 'admin1234'):
        import time
        session['admin_logged_in'] = True
        session['admin_auth_issued_at'] = time.time()
        return jsonify({'success': True})

    return jsonify({'error': 'invalid_password', 'message': 'پسورد اشتباه است'}), 401


@admin_bp.route('/logout', methods=['POST'])
def admin_logout():
    session.pop('admin_logged_in', None)
    return jsonify({'success': True})


@admin_bp.route('/status')
def admin_status():
    return jsonify({'logged_in': bool(session.get('admin_logged_in'))})


@admin_bp.route('/users')
@admin_required
def list_users():
    page = request.args.get('page', 1, type=int)
    per_page = request.args.get('per_page', 50, type=int)

    users_page = (User.query
                  .order_by(User.created_at.desc())
                  .paginate(page=page, per_page=per_page, error_out=False))

    return jsonify({
        'users': [u.to_dict() for u in users_page.items],
        'total': users_page.total,
        'pages': users_page.pages,
        'current_page': page
    })


@admin_bp.route('/conversations')
@admin_required
def list_conversations():
    page = request.args.get('page', 1, type=int)
    per_page = request.args.get('per_page', 50, type=int)

    convs_page = (Conversation.query
                  .order_by(Conversation.started_at.desc())
                  .paginate(page=page, per_page=per_page, error_out=False))

    result = []
    for c in convs_page.items:
        d = c.to_dict()
        if c.user:
            d['phone_number'] = c.user.phone_number
        result.append(d)

    return jsonify({
        'conversations': result,
        'total': convs_page.total,
        'pages': convs_page.pages,
        'current_page': page
    })


@admin_bp.route('/messages/<int:conv_id>')
@admin_required
def get_messages(conv_id):
    conv = Conversation.query.get_or_404(conv_id)
    user_phone = conv.user.phone_number if conv.user else 'ناشناس'

    return jsonify({
        'conversation_id': conv_id,
        'user_phone': user_phone,
        'started_at': conv.started_at.isoformat() if conv.started_at else None,
        'messages': [m.to_dict() for m in conv.messages]
    })


@admin_bp.route('/search')
@admin_required
def search_messages():
    q = request.args.get('q', '').strip()
    if len(q) < 2:
        return jsonify({'results': [], 'message': 'حداقل ۲ کاراکتر وارد کنید'})

    results = (Message.query
               .filter(Message.content.contains(q))
               .order_by(Message.timestamp.desc())
               .limit(100)
               .all())

    output = []
    for m in results:
        d = m.to_dict()
        if m.conversation and m.conversation.user:
            d['phone_number'] = m.conversation.user.phone_number
        output.append(d)

    return jsonify({'results': output, 'count': len(output)})


@admin_bp.route('/contact-requests')
@admin_required
def list_contact_requests():
    requests_list = (ContactRequest.query
                     .order_by(ContactRequest.created_at.desc())
                     .all())
    return jsonify({'requests': [r.to_dict() for r in requests_list]})


@admin_bp.route('/contact-requests/<int:req_id>/read', methods=['POST'])
@admin_required
def mark_contact_read(req_id):
    cr = ContactRequest.query.get_or_404(req_id)
    cr.is_read = True
    db.session.commit()
    return jsonify({'success': True})


@admin_bp.route('/contact-requests/<int:req_id>/delete', methods=['POST'])
@admin_required
def delete_contact_request(req_id):
    cr = ContactRequest.query.get_or_404(req_id)
    db.session.delete(cr)
    db.session.commit()
    return jsonify({'success': True})


@admin_bp.route('/assessment-leads')
@admin_required
def list_assessment_leads():
    leads = AssessmentLead.query.order_by(AssessmentLead.created_at.desc()).all()
    result = []
    for lead in leads:
        d = lead.to_dict()
        types = []
        if Assessment.query.filter_by(lead_id=lead.id).first() is not None:
            types.append('فرسودگی شغلی')
        if WaaqAssessment.query.filter_by(lead_id=lead.id).first() is not None:
            types.append('انعطاف پذیری')
        if CareerKnotAssessment.query.filter_by(lead_id=lead.id).first() is not None:
            types.append('گره کور شغلی')
        d['assessment_types'] = types
        d['has_assessment'] = len(types) > 0
        result.append(d)
    return jsonify({'leads': result})


@admin_bp.route('/assessment-leads/<int:lead_id>/delete', methods=['POST'])
@admin_required
def delete_assessment_lead(lead_id):
    lead = AssessmentLead.query.get_or_404(lead_id)
    db.session.delete(lead)
    db.session.commit()
    return jsonify({'success': True})


@admin_bp.route('/webinar-registrations')
@admin_required
def list_webinar_registrations():
    regs = (WebinarRegistration.query
            .order_by(WebinarRegistration.created_at.desc())
            .all())
    return jsonify({'registrations': [r.to_dict() for r in regs]})


@admin_bp.route('/webinar-registrations/<int:reg_id>/delete', methods=['POST'])
@admin_required
def delete_webinar_registration(reg_id):
    reg = WebinarRegistration.query.get_or_404(reg_id)
    db.session.delete(reg)
    db.session.commit()
    return jsonify({'success': True})


@admin_bp.route('/knot-webinar-registrations')
@admin_required
def list_knot_webinar_registrations():
    regs = (WebinarRegistration.query
            .filter_by(webinar_title=KNOT_WEBINAR_TITLE)
            .order_by(WebinarRegistration.created_at.desc())
            .all())
    return jsonify({'registrations': [r.to_dict() for r in regs]})


@admin_bp.route('/knot-webinar-registrations/<int:reg_id>/delete', methods=['POST'])
@admin_required
def delete_knot_webinar_registration(reg_id):
    reg = WebinarRegistration.query.get_or_404(reg_id)
    db.session.delete(reg)
    db.session.commit()
    return jsonify({'success': True})


@admin_bp.route('/landing-webinar-registrations')
@admin_required
def list_landing_webinar_registrations():
    """ثبت‌نام‌هایی که منحصراً از طریق لندینگ پیج وبینار صورت گرفته‌اند."""
    regs = (WebinarRegistration.query
            .filter_by(source='landing')
            .order_by(WebinarRegistration.created_at.desc())
            .all())
    return jsonify({'registrations': [r.to_dict() for r in regs]})


@admin_bp.route('/landing-webinar-registrations/<int:reg_id>/delete', methods=['POST'])
@admin_required
def delete_landing_webinar_registration(reg_id):
    reg = WebinarRegistration.query.get_or_404(reg_id)
    db.session.delete(reg)
    db.session.commit()
    return jsonify({'success': True})


@admin_bp.route('/assessment-leads/<int:lead_id>/details')
@admin_required
def assessment_lead_details(lead_id):
    lead = AssessmentLead.query.get_or_404(lead_id)

    knot_option_fa = {1: 'الف', 2: 'ب', 3: 'ج', 4: 'د'}
    assessments_out = []

    # Career Knot
    for ka in (CareerKnotAssessment.query.filter_by(lead_id=lead.id)
               .order_by(CareerKnotAssessment.completed_at.desc()).all()):
        import json as json_lib
        comments = []
        if ka.comments:
            try:
                comments = json_lib.loads(ka.comments)
            except (ValueError, TypeError):
                comments = []
        assessments_out.append({
            'type': 'گره کور شغلی',
            'completed_at': ka.completed_at.isoformat() if ka.completed_at else None,
            'answers': [
                {
                    'question': i + 1,
                    'answer': knot_option_fa.get(v, ''),
                    'comment': (comments[i] if isinstance(comments, list) and i < len(comments) else None)
                }
                for i, v in enumerate([ka.item1, ka.item2, ka.item3, ka.item4,
                                       ka.item5, ka.item6, ka.item7, ka.item8])
            ],
            'summary': f"{ka.profile_title} — الف:{ka.count_a} ب:{ka.count_b} ج:{ka.count_c} د:{ka.count_d}"
        })

    # MBI (فرسودگی شغلی)
    level_map = {'low': 'پایین', 'moderate': 'متوسط', 'high': 'بالا'}
    for a in Assessment.query.filter_by(lead_id=lead.id).order_by(Assessment.completed_at.desc()).all():
        assessments_out.append({
            'type': 'فرسودگی شغلی',
            'completed_at': a.completed_at.isoformat() if a.completed_at else None,
            'answers': [],
            'summary': (f"خستگی عاطفی: {a.emotional_exhaustion} ({level_map.get(a.ee_level, '')}) | "
                        f"مسخ شخصیت: {a.depersonalization} ({level_map.get(a.dp_level, '')}) | "
                        f"کفایت فردی: {a.personal_accomplishment} ({level_map.get(a.pa_level, '')})")
        })

    # WAAQ (انعطاف پذیری)
    for wa in WaaqAssessment.query.filter_by(lead_id=lead.id).order_by(WaaqAssessment.completed_at.desc()).all():
        assessments_out.append({
            'type': 'انعطاف پذیری',
            'completed_at': wa.completed_at.isoformat() if wa.completed_at else None,
            'answers': [
                {'question': i + 1, 'answer': v, 'comment': None}
                for i, v in enumerate([wa.item1, wa.item2, wa.item3, wa.item4, wa.item5, wa.item6, wa.item7])
            ],
            'summary': f"نمره کل: {wa.total_score} — سطح: {level_map.get(wa.level, '')}"
        })

    return jsonify({'lead': lead.to_dict(), 'assessments': assessments_out})


@admin_bp.route('/orders')
@admin_required
def list_orders():
    orders = Order.query.order_by(Order.created_at.desc()).all()
    return jsonify({'orders': [o.to_dict() for o in orders]})


@admin_bp.route('/users/<int:user_id>/orders')
@admin_required
def list_user_orders(user_id):
    """سفارشات تکمیل‌شده یک کاربر خاص — دسترسی ادمین به پنل کاربر."""
    orders = (Order.query
              .filter_by(user_id=user_id, status='completed')
              .order_by(Order.completed_at.desc())
              .all())
    return jsonify({'user_id': user_id, 'orders': [o.to_dict() for o in orders]})


@admin_bp.route('/orders/<int:order_id>/set-key', methods=['POST'])
@admin_required
def set_order_key(order_id):
    data = request.get_json()
    key = (data or {}).get('key', '').strip()
    if not key:
        return jsonify({'error': 'key_required'}), 400

    order = Order.query.get_or_404(order_id)
    order.spotplayer_key = key
    order.status = 'completed'
    order.completed_at = iran_now()
    db.session.commit()

    # Send key via SMS
    try:
        from sms_service import send_sms
        phone = order.user.phone_number
        product_name = order.product.name
        msg = f'سای‌باما: کلید دسترسی به {product_name} شما:\n{key}\nبا موفقیت فعال شد.'
        send_sms(phone, msg)
    except Exception:
        pass

    return jsonify({'success': True})


@admin_bp.route('/orders/<int:order_id>/cancel', methods=['POST'])
@admin_required
def cancel_order(order_id):
    order = Order.query.get_or_404(order_id)
    order.status = 'cancelled'
    db.session.commit()
    return jsonify({'success': True})


@admin_bp.route('/orders/<int:order_id>/mark-shipped', methods=['POST'])
@admin_required
def mark_shipped(order_id):
    """برای محصولات فیزیکی (مثل کتاب): علامت ارسال شد + اطلاع‌رسانی پیامکی (بدون کلید)."""
    order = Order.query.get_or_404(order_id)
    order.status = 'completed'
    order.completed_at = iran_now()
    db.session.commit()

    try:
        from sms_service import send_sms
        phone = order.user.phone_number if order.user else order.customer_phone
        product_name = order.product.name if order.product else 'سفارش'
        msg = f'سای‌باما: سفارش «{product_name}» شما با موفقیت ثبت و ارسال شد. با تشکر از خرید شما.'
        send_sms(phone, msg)
    except Exception:
        pass

    return jsonify({'success': True})


@admin_bp.route('/products')
@admin_required
def list_products():
    products = Product.query.all()
    return jsonify({'products': [p.to_dict() for p in products]})


@admin_bp.route('/products/<int:product_id>', methods=['POST'])
@admin_required
def update_product(product_id):
    data = request.get_json()
    product = Product.query.get_or_404(product_id)
    if 'name' in data:
        product.name = data['name']
    if 'description' in data:
        product.description = data['description']
    if 'price' in data:
        product.price = int(data['price'])
    if 'is_active' in data:
        product.is_active = bool(data['is_active'])
    db.session.commit()
    return jsonify({'success': True, 'product': product.to_dict()})


@admin_bp.route('/stats')
@admin_required
def get_stats():
    pending = Order.query.filter(Order.status.in_(['pending_payment', 'pending_key'])).count()
    return jsonify({
        'total_users': User.query.count(),
        'verified_users': User.query.filter_by(is_verified=True).count(),
        'total_conversations': Conversation.query.count(),
        'total_messages': Message.query.count(),
        'total_assessments': Assessment.query.count(),
        'total_waaq_assessments': WaaqAssessment.query.count(),
        'total_knot_assessments': CareerKnotAssessment.query.count(),
        'total_webinar_registrations': WebinarRegistration.query.count(),
        'total_contact_requests': ContactRequest.query.count(),
        'unread_contact_requests': ContactRequest.query.filter_by(is_read=False).count(),
        'pending_orders': pending,
        'total_career_intakes': CareerIntake.query.count(),
        'total_org_intakes': OrgIntake.query.count(),
    })


@admin_bp.route('/career-intakes')
@admin_required
def list_career_intakes():
    items = (CareerIntake.query
             .order_by(CareerIntake.created_at.desc())
             .all())
    return jsonify({'intakes': [
        {'id': i.id, 'name': i.q1_name, 'job': i.q4_job,
         'created_at': i.created_at.isoformat() if i.created_at else None}
        for i in items
    ]})


@admin_bp.route('/career-intakes/<int:item_id>')
@admin_required
def career_intake_detail(item_id):
    item = CareerIntake.query.get_or_404(item_id)
    return jsonify({'intake': item.to_dict()})


@admin_bp.route('/career-intakes/<int:item_id>/delete', methods=['POST'])
@admin_required
def delete_career_intake(item_id):
    item = CareerIntake.query.get_or_404(item_id)
    db.session.delete(item)
    db.session.commit()
    return jsonify({'success': True})


@admin_bp.route('/org-intakes')
@admin_required
def list_org_intakes():
    items = (OrgIntake.query
             .order_by(OrgIntake.created_at.desc())
             .all())
    return jsonify({'intakes': [
        {'id': i.id, 'org': i.q1_org, 'filler': i.q2_filler,
         'created_at': i.created_at.isoformat() if i.created_at else None}
        for i in items
    ]})


@admin_bp.route('/org-intakes/<int:item_id>')
@admin_required
def org_intake_detail(item_id):
    item = OrgIntake.query.get_or_404(item_id)
    return jsonify({'intake': item.to_dict()})


@admin_bp.route('/org-intakes/<int:item_id>/delete', methods=['POST'])
@admin_required
def delete_org_intake(item_id):
    item = OrgIntake.query.get_or_404(item_id)
    db.session.delete(item)
    db.session.commit()
    return jsonify({'success': True})


@admin_bp.route('/export')
@admin_required
def export_excel():
    wb = Workbook()

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(fill_type='solid', fgColor='1D5FA6')
    center_align = Alignment(horizontal='center', vertical='center')

    # Sheet 1: Users
    ws_users = wb.active
    ws_users.title = 'کاربران'
    headers_u = ['شناسه', 'شماره موبایل', 'تأیید شده', 'تاریخ ثبت‌نام', 'تعداد گفتگو']
    ws_users.append(headers_u)
    for cell in ws_users[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for u in User.query.order_by(User.created_at).all():
        ws_users.append([
            u.id,
            u.phone_number,
            'بله' if u.is_verified else 'خیر',
            str(u.created_at)[:19] if u.created_at else '',
            len(u.conversations)
        ])

    ws_users.column_dimensions['B'].width = 20
    ws_users.column_dimensions['D'].width = 22

    # Sheet 2: Messages
    ws_msgs = wb.create_sheet('گفتگوها')
    headers_m = ['شناسه گفتگو', 'شماره موبایل', 'نقش', 'محتوا', 'زمان']
    ws_msgs.append(headers_m)
    for cell in ws_msgs[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    messages = (Message.query
                .join(Conversation)
                .join(User)
                .order_by(Conversation.id, Message.timestamp)
                .all())

    for m in messages:
        phone = m.conversation.user.phone_number if (m.conversation and m.conversation.user) else ''
        role_fa = 'کاربر' if m.role == 'user' else 'سای‌باما'
        ws_msgs.append([
            m.conversation_id,
            phone,
            role_fa,
            m.content,
            str(m.timestamp)[:19] if m.timestamp else ''
        ])

    ws_msgs.column_dimensions['B'].width = 18
    ws_msgs.column_dimensions['D'].width = 60
    ws_msgs.column_dimensions['E'].width = 22

    # Sheet 3: Assessments
    ws_assess = wb.create_sheet('پرسشنامه‌ها')
    headers_a = ['شناسه', 'شماره موبایل', 'خستگی عاطفی', 'سطح EE', 'مسخ شخصیت', 'سطح DP', 'کفایت فردی', 'سطح PA', 'تاریخ']
    ws_assess.append(headers_a)
    for cell in ws_assess[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for a in Assessment.query.all():
        phone = a.user.phone_number if a.user else ''
        level_map = {'low': 'پایین', 'moderate': 'متوسط', 'high': 'بالا'}
        ws_assess.append([
            a.id, phone,
            a.emotional_exhaustion, level_map.get(a.ee_level, ''),
            a.depersonalization, level_map.get(a.dp_level, ''),
            a.personal_accomplishment, level_map.get(a.pa_level, ''),
            str(a.completed_at)[:19] if a.completed_at else ''
        ])

    # Sheet 4: Contact Requests
    ws_contacts = wb.create_sheet('درخواست مشاوره')
    headers_c = ['شناسه', 'نام', 'شماره موبایل', 'پیام', 'تاریخ', 'خوانده شده']
    ws_contacts.append(headers_c)
    for cell in ws_contacts[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for cr in ContactRequest.query.order_by(ContactRequest.created_at.desc()).all():
        ws_contacts.append([
            cr.id,
            cr.name,
            cr.phone,
            cr.message or '',
            str(cr.created_at)[:19] if cr.created_at else '',
            'بله' if cr.is_read else 'خیر'
        ])

    ws_contacts.column_dimensions['B'].width = 20
    ws_contacts.column_dimensions['C'].width = 18
    ws_contacts.column_dimensions['D'].width = 40
    ws_contacts.column_dimensions['E'].width = 22

    # Sheet 5: Assessment Leads
    ws_leads = wb.create_sheet('لیدهای پرسشنامه')
    headers_l = ['شناسه', 'نام', 'ایمیل', 'تلفن', 'تاریخ', 'تکمیل پرسشنامه']
    ws_leads.append(headers_l)
    for cell in ws_leads[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for lead in AssessmentLead.query.order_by(AssessmentLead.created_at.desc()).all():
        types = []
        if Assessment.query.filter_by(lead_id=lead.id).first() is not None:
            types.append('فرسودگی شغلی')
        if WaaqAssessment.query.filter_by(lead_id=lead.id).first() is not None:
            types.append('انعطاف پذیری')
        if CareerKnotAssessment.query.filter_by(lead_id=lead.id).first() is not None:
            types.append('گره کور شغلی')
        ws_leads.append([
            lead.id,
            lead.name,
            lead.email,
            lead.phone,
            str(lead.created_at)[:19] if lead.created_at else '',
            ' + '.join(types) if types else 'خیر'
        ])

    ws_leads.column_dimensions['B'].width = 20
    ws_leads.column_dimensions['C'].width = 26
    ws_leads.column_dimensions['D'].width = 18
    ws_leads.column_dimensions['E'].width = 22

    # Sheet 6: Orders
    ws_orders = wb.create_sheet('سفارش‌ها')
    headers_o = ['شناسه', 'شماره موبایل', 'محصول', 'مبلغ', 'روش پرداخت', 'وضعیت', 'تاریخ ثبت']
    ws_orders.append(headers_o)
    for cell in ws_orders[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for o in Order.query.order_by(Order.created_at.desc()).all():
        phone = o.user.phone_number if o.user else ''
        ws_orders.append([
            o.id,
            phone,
            o.product.name if o.product else '',
            o.amount,
            o.payment_method,
            o.status,
            str(o.created_at)[:19] if o.created_at else ''
        ])

    ws_orders.column_dimensions['B'].width = 18
    ws_orders.column_dimensions['C'].width = 26
    ws_orders.column_dimensions['G'].width = 22

    # Sheet 7: WAAQ Assessments
    ws_waaq = wb.create_sheet('پرسشنامه WAAQ')
    headers_w = ['شناسه', 'شماره موبایل', 'سوال ۱', 'سوال ۲', 'سوال ۳', 'سوال ۴', 'سوال ۵', 'سوال ۶', 'سوال ۷', 'نمره کل', 'سطح', 'تاریخ']
    ws_waaq.append(headers_w)
    for cell in ws_waaq[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for wa in WaaqAssessment.query.order_by(WaaqAssessment.completed_at.desc()).all():
        phone = wa.user.phone_number if wa.user else ''
        level_map = {'low': 'پایین', 'moderate': 'متوسط', 'high': 'بالا'}
        ws_waaq.append([
            wa.id, phone,
            wa.item1, wa.item2, wa.item3, wa.item4, wa.item5, wa.item6, wa.item7,
            wa.total_score, level_map.get(wa.level, ''),
            str(wa.completed_at)[:19] if wa.completed_at else ''
        ])

    ws_waaq.column_dimensions['B'].width = 18
    ws_waaq.column_dimensions['K'].width = 22

    # Sheet 8: Career Knot Assessments (پرسشنامه ریشه‌یابی گره کور شغلی)
    ws_knot = wb.create_sheet('پرسشنامه گره کور شغلی')
    headers_k = ['شناسه', 'نام', 'ایمیل', 'شماره تماس',
                 'سوال ۱', 'سوال ۲', 'سوال ۳', 'سوال ۴', 'سوال ۵', 'سوال ۶', 'سوال ۷', 'سوال ۸',
                 'تعداد الف', 'تعداد ب', 'تعداد ج', 'تعداد د', 'نمره کل (T)', 'بخش تفسیر', 'پروفایل', 'توضیحات', 'تاریخ']
    ws_knot.append(headers_k)
    for cell in ws_knot[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    knot_option_fa = {1: 'الف', 2: 'ب', 3: 'ج', 4: 'د'}
    for ka in CareerKnotAssessment.query.order_by(CareerKnotAssessment.completed_at.desc()).all():
        if ka.user:
            name, email, phone = '', '', ka.user.phone_number
        elif ka.lead:
            name, email, phone = ka.lead.name, ka.lead.email, ka.lead.phone
        else:
            name, email, phone = '', '', ''

        comments_text = ''
        if ka.comments:
            try:
                import json as _json
                items = _json.loads(ka.comments)
                parts = [f'س{i + 1}: {c}' for i, c in enumerate(items) if c]
                comments_text = ' | '.join(parts)
            except (ValueError, TypeError):
                comments_text = str(ka.comments)

        ws_knot.append([
            ka.id, name, email, phone,
            knot_option_fa.get(ka.item1, ''), knot_option_fa.get(ka.item2, ''),
            knot_option_fa.get(ka.item3, ''), knot_option_fa.get(ka.item4, ''),
            knot_option_fa.get(ka.item5, ''), knot_option_fa.get(ka.item6, ''),
            knot_option_fa.get(ka.item7, ''), knot_option_fa.get(ka.item8, ''),
            ka.count_a, ka.count_b, ka.count_c, ka.count_d,
            ka.total_score,
            f'بخش {ka.section}',
            ka.profile_title,
            comments_text,
            str(ka.completed_at)[:19] if ka.completed_at else ''
        ])

    ws_knot.column_dimensions['B'].width = 18
    ws_knot.column_dimensions['C'].width = 24
    ws_knot.column_dimensions['D'].width = 16
    ws_knot.column_dimensions['T'].width = 40
    ws_knot.column_dimensions['U'].width = 22

    # Sheet 9: Webinar Registrations (ثبت‌نام وبینارها)
    ws_web = wb.create_sheet('ثبت‌نام وبینار')
    headers_web = ['شناسه', 'نام و نام خانوادگی', 'ایمیل', 'شماره تماس',
                   'وبینار', 'تاریخ برگزاری', 'تاریخ ثبت‌نام', 'پیامک ارسال شد']
    ws_web.append(headers_web)
    for cell in ws_web[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for r in WebinarRegistration.query.order_by(WebinarRegistration.created_at.desc()).all():
        ws_web.append([
            r.id,
            r.name or '',
            r.email or '',
            r.phone,
            r.webinar_title,
            r.webinar_date or '',
            str(r.created_at)[:19] if r.created_at else '',
            'بله' if r.sms_sent else 'خیر'
        ])

    ws_web.column_dimensions['B'].width = 20
    ws_web.column_dimensions['C'].width = 26
    ws_web.column_dimensions['D'].width = 16
    ws_web.column_dimensions['E'].width = 32
    ws_web.column_dimensions['F'].width = 18
    ws_web.column_dimensions['G'].width = 22

    # Sheet 10: Career Intake Forms (فرم پذیرش کوچینگ شغلی)
    ws_career = wb.create_sheet('فرم پذیرش کوچینگ شغلی')
    headers_car = ['شناسه', 'نام و نام خانوادگی', 'سن / وضعیت تأهل', 'شماره تماس', 'ایمیل',
                   'شغل فعلی', 'شهر', 'میزان درآمد', 'گره کور شغلی', 'نتیجه ایده‌آل',
                   'اقدامات قبلی', 'سابقه درمان', 'میزان تعهد', 'اولویت سرمایه‌گذاری', 'تاریخ ثبت']
    ws_career.append(headers_car)
    for cell in ws_career[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for i in CareerIntake.query.order_by(CareerIntake.created_at.desc()).all():
        ws_career.append([
            i.id, i.q1_name, i.q2_marital, i.q3_phone or i.q3_contact or '', i.q3_email or '',
            i.q4_job, i.q5_city, i.q6_income,
            i.q7_knot, i.q8_ideal, i.q9_actions,
            i.q10_treatment, i.q11_commitment, i.q12_priority,
            str(i.created_at)[:19] if i.created_at else ''
        ])

    ws_career.column_dimensions['B'].width = 24
    ws_career.column_dimensions['C'].width = 20
    ws_career.column_dimensions['D'].width = 26
    ws_career.column_dimensions['H'].width = 60
    ws_career.column_dimensions['I'].width = 60
    ws_career.column_dimensions['J'].width = 60
    ws_career.column_dimensions['N'].width = 22

    # Sheet 11: Org Intake Forms (فرم پذیرش کوچینگ سازمانی)
    ws_org = wb.create_sheet('فرم پذیرش کوچینگ سازمانی')
    headers_org = ['شناسه', 'نام سازمان', 'نام و سمت تکمیل‌کننده', 'شماره تماس', 'ایمیل', 'زمینه فعالیت',
                   'تعداد پرسنل', 'گردش مالی سالانه', 'بودجه مصوب', 'تمرکز کوچینگ',
                   'چالش‌های رفتاری', 'ریشه چالش‌ها', 'KPI', 'آمادگی هیئت‌مدیره',
                   'افق زمانی', 'تاریخ ثبت']
    ws_org.append(headers_org)
    for cell in ws_org[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for i in OrgIntake.query.order_by(OrgIntake.created_at.desc()).all():
        ws_org.append([
            i.id, i.q1_org, i.q2_filler, i.q_contact_phone or '', i.q_contact_email or '', i.q3_industry,
            i.q4_staff, i.q5_turnover, i.q6_budget, i.q7_focus,
            i.q8_symptoms, i.q9_root, i.q10_kpi, i.q11_readiness,
            i.q12_horizon,
            str(i.created_at)[:19] if i.created_at else ''
        ])

    ws_org.column_dimensions['B'].width = 28
    ws_org.column_dimensions['C'].width = 26
    ws_org.column_dimensions['G'].width = 22
    ws_org.column_dimensions['I'].width = 60
    ws_org.column_dimensions['J'].width = 60
    ws_org.column_dimensions['K'].width = 60
    ws_org.column_dimensions['M'].width = 22

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='psybama_export.xlsx'
    )
